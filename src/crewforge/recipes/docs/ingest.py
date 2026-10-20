"""Parse documents and split them into overlapping chunks (no LLM calls)."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ...core.errors import ConfigError
from .settings import DocsSettings

_DOCS_EXTRA = "pip install 'crewforge[docs]'"


@dataclass
class TextChunk:
    """A chunk of text from a document."""

    chunk_id: str
    document_id: str
    text: str
    start_offset: int
    end_offset: int
    page_number: int | None = None
    section: str = ""
    word_count: int = 0
    char_count: int = 0

    def __post_init__(self) -> None:
        self.word_count = len(self.text.split())
        self.char_count = len(self.text)


@dataclass
class Document:
    """Parsed document with metadata."""

    doc_id: str
    filename: str
    content: str
    file_type: str
    page_count: int = 0
    chunks: list[TextChunk] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def total_words(self) -> int:
        return sum(c.word_count for c in self.chunks)


class IngestAgent:
    """Read a file, detect its type and split it into chunks. Deterministic."""

    name = "ingester"

    TYPE_MAP = {
        ".txt": "text",
        ".md": "markdown",
        ".pdf": "pdf",
        ".docx": "docx",
        ".html": "html",
        ".htm": "html",
        ".json": "json",
        ".csv": "csv",
        ".rst": "rst",
    }

    def __init__(self, settings: DocsSettings | None = None) -> None:
        self.settings = settings or DocsSettings()
        self.chunk_size = self.settings.chunk_size
        self.chunk_overlap = self.settings.chunk_overlap

    def _detect_type(self, path: Path) -> str:
        return self.TYPE_MAP.get(path.suffix.lower(), "text")

    def _read_file(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._read_pdf(path)
        if suffix == ".docx":
            return self._read_docx(path)
        if suffix in (".html", ".htm"):
            return self._read_html(path)
        return path.read_text(encoding="utf-8", errors="replace")

    def _read_pdf(self, path: Path) -> str:
        """Extract text with PyMuPDF (optional ``docs`` extra).

        Without it we refuse instead of reading the binary as text: a PDF's raw
        bytes "ingest" fine and then every answer is grounded in garbage.
        """
        try:
            import pymupdf
        except ImportError:
            try:
                import fitz as pymupdf  # older PyMuPDF releases
            except ImportError as exc:
                raise ConfigError(
                    f"reading PDF files needs PyMuPDF: {_DOCS_EXTRA} ({path.name})"
                ) from exc
        doc = pymupdf.open(str(path))
        try:
            return "\n\n".join(page.get_text() for page in doc)
        finally:
            doc.close()

    def _read_docx(self, path: Path) -> str:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise ConfigError(
                f"reading DOCX files needs python-docx: {_DOCS_EXTRA} ({path.name})"
            ) from exc
        doc = DocxDocument(str(path))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _read_html(self, path: Path) -> str:
        raw = path.read_text(encoding="utf-8", errors="replace")
        raw = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", raw)
        return re.sub(r"<[^>]+>", " ", raw)

    def _generate_doc_id(self, content: str, filename: str) -> str:
        """Content-addressed id: the same text gets the same id whatever its filename."""
        h = hashlib.sha256(content[:4096].encode()).hexdigest()[:12]
        return f"doc-{h}"

    def _split_into_chunks(self, text: str, doc_id: str, start_index: int = 0) -> list[TextChunk]:
        """Split text on sentence boundaries into overlapping chunks."""
        chunks: list[TextChunk] = []
        sentences = re.split(r"(?<=[.!?])\s+", text)
        current = ""
        offset = 0
        chunk_idx = start_index

        for sentence in sentences:
            if len(current) + len(sentence) + 1 > self.chunk_size and current:
                chunks.append(
                    TextChunk(
                        chunk_id=f"{doc_id}-chunk-{chunk_idx:04d}",
                        document_id=doc_id,
                        text=current.strip(),
                        start_offset=offset,
                        end_offset=offset + len(current),
                    )
                )
                overlap_text = current[-self.chunk_overlap :] if self.chunk_overlap else ""
                offset += len(current) - len(overlap_text)
                current = overlap_text + " " + sentence
                chunk_idx += 1
            else:
                current = (current + " " + sentence).strip()

        if current.strip():
            chunks.append(
                TextChunk(
                    chunk_id=f"{doc_id}-chunk-{chunk_idx:04d}",
                    document_id=doc_id,
                    text=current.strip(),
                    start_offset=offset,
                    end_offset=offset + len(current),
                )
            )
        return chunks

    def _extract_sections(self, text: str) -> list[tuple[str, str]]:
        """Split Markdown/RST-style text on ``#`` headers into (header, body) pairs."""
        sections: list[tuple[str, str]] = []
        matches = list(re.finditer(r"^(#{1,6})\s+(.+)$", text, re.MULTILINE))
        if not matches:
            return [("main", text)]
        for i, match in enumerate(matches):
            start = match.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            content = text[start:end].strip()
            if content:
                sections.append((match.group(2).strip(), content))
        return sections

    async def run(self, path: str | Path, *, extract_sections: bool = True) -> Document:
        """Ingest one document: read, parse, chunk."""
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"Document not found: {p}")

        content = self._read_file(p)
        doc_id = self._generate_doc_id(content, p.name)
        file_type = self._detect_type(p)
        doc = Document(
            doc_id=doc_id,
            filename=p.name,
            content=content,
            file_type=file_type,
            metadata={"source_path": str(p.resolve()), "size_bytes": p.stat().st_size},
        )

        if extract_sections and file_type in ("markdown", "rst"):
            all_chunks: list[TextChunk] = []
            for section_name, section_text in self._extract_sections(content):
                # Continue numbering across sections so chunk ids stay unique.
                section_chunks = self._split_into_chunks(section_text, doc_id, len(all_chunks))
                for c in section_chunks:
                    c.section = section_name
                all_chunks.extend(section_chunks)
            doc.chunks = all_chunks
        else:
            doc.chunks = self._split_into_chunks(content, doc_id)
        return doc

    async def run_batch(self, paths: list[str | Path]) -> list[Document]:
        return [await self.run(p) for p in paths]
